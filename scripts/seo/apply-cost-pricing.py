"""Set every catalog price to Sanitube cost x a multiplier.

    python3 scripts/seo/apply-cost-pricing.py --mult 1.8                 # dry run: show every change, write nothing
    python3 scripts/seo/apply-cost-pricing.py --mult 1.8 --apply         # write, after saving a backup of the old prices
    python3 scripts/seo/apply-cost-pricing.py --revert pricing-backup-<time>.json   # put the old prices back
    python3 scripts/seo/apply-cost-pricing.py --mult 1.8 --ferguson Cardinal-Cooling-Pricing.xlsx   # cap at Ferguson

    Cowbird (same catalog, same cost file):
    python3 scripts/seo/apply-cost-pricing.py --mult 1.75 \
        --backend https://cowbird-depot-production.up.railway.app \
        --env ~/colibri-repos/cowbirddepot.com/backend/.env

Cost comes from Sanitube's price file ("Sanitube A Pricing 02-02-2026.csv", column "Correct Price", one row per
SKU). Price = round(cost x mult, 2) in USD. Tube (AI270/A270 SKUs) is skipped: it's priced per foot from its own
list by apply-tube-pricing.py. SKUs with a $0 cost, rows whose Notes aren't "OK", and site SKUs missing from the
file are left alone and listed so someone can decide what to do with them.

With --ferguson (the owner's sheet: SKU in column A, Ferguson price in column F), a SKU whose cost x mult is above
Ferguson drops to --under (10%) below Ferguson, but never under --min-margin (20%) gross margin. Rare sizes where
Ferguson sells below our cost stay above Ferguson at that floor. Always pass it, or a plain re-run raises them again.

Every price on a variant is sent back on update (Medusa replaces a variant's whole price set), so prices in other
currencies or with quantity rules survive untouched. Admin credentials come from backend/.env like the other
scripts here. Idempotent: re-running with the same multiplier changes nothing.
"""
import argparse, csv, json, math, os, pathlib, sys, time, urllib.request

HERE = pathlib.Path(__file__).parent; ROOT = HERE.parent.parent
DEFAULT_CSV = ROOT.parent / "cardinalcoolingsystems.com" / "Sanitube A Pricing 02-02-2026.csv"


def env(k, path):
    if os.environ.get(k):
        return os.environ[k]
    files = [pathlib.Path(path).expanduser()] if path else [ROOT / "backend/.env", ROOT.parent / "cardinalcoolingsystems.com/backend/.env"]
    for f in files:
        if f.exists():
            for line in f.read_text().splitlines():
                if line.startswith(k + "="):
                    return line.split("=", 1)[1].strip().strip('"')
    sys.exit(f"missing {k}")


class Admin:
    def __init__(self, base, env_path):
        self.base = base.rstrip("/")
        self.tok = self.call("/auth/user/emailpass", {"email": env("MEDUSA_ADMIN_EMAIL", env_path), "password": env("MEDUSA_ADMIN_PASSWORD", env_path)})["token"]

    def call(self, path, body=None):
        h = {"Content-Type": "application/json", "User-Agent": "cardinal-pricing/1"}
        if getattr(self, "tok", None):
            h["Authorization"] = f"Bearer {self.tok}"
        req = urllib.request.Request(self.base + path, data=json.dumps(body).encode() if body is not None else None, headers=h,
                                     method="POST" if body is not None else "GET")
        for attempt in range(4):
            try:
                return json.load(urllib.request.urlopen(req, timeout=120))
            except urllib.error.HTTPError as e:
                if e.code >= 500 and attempt < 3:
                    time.sleep(3 * (attempt + 1)); continue
                sys.exit(f"{path} -> {e.code}: {e.read().decode()[:400]}")

    def products(self):
        out, off = [], 0
        while True:
            r = self.call(f"/admin/products?limit=100&offset={off}&fields=id,title,handle,*variants,variants.id,variants.sku,variants.title,*variants.prices")
            out += r["products"]; off += 100
            if off >= r["count"]:
                return out


def is_tube(sku):
    return sku.startswith(("AI270", "A270"))


def ferguson_prices(path):
    from openpyxl import load_workbook
    out = {}
    for r in load_workbook(path, read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True):
        if r[0] and isinstance(r[5], (int, float)) and r[5] > 0:
            out[str(r[0]).strip()] = float(r[5])
    return out


def target_price(cost, mult, ferg, under, min_margin):
    """cost x mult, capped at `under` below Ferguson, never below `min_margin` gross margin."""
    price = round(cost * mult, 2)
    if ferg is None or price <= ferg:
        return price
    floor = math.ceil(cost / (1 - min_margin) * 100 - 1e-6) / 100
    return min(price, max(round(ferg * (1 - under), 2), floor))


def price_payload(prices, usd_amount=None):
    """All of a variant's prices, with the plain USD one (no rules) set to usd_amount."""
    rows = []
    for p in prices:
        row = {"id": p["id"], "currency_code": p["currency_code"], "amount": p["amount"]}
        if usd_amount is not None and p is plain_usd(prices):
            row["amount"] = usd_amount
        if p.get("min_quantity") is not None: row["min_quantity"] = p["min_quantity"]
        if p.get("max_quantity") is not None: row["max_quantity"] = p["max_quantity"]
        rows.append(row)
    return rows


def plain_usd(prices):
    usd = [p for p in prices if p["currency_code"] == "usd" and not p.get("min_quantity") and not p.get("max_quantity") and not (p.get("rules") or {})]
    return usd[0] if usd else None


def send(admin, by_product):
    done = 0
    for pid, updates in by_product.items():
        for i in range(0, len(updates), 50):
            admin.call(f"/admin/products/{pid}/variants/batch", {"update": updates[i:i + 50]})
            done += len(updates[i:i + 50])
    return done


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--mult", type=float)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--revert")
    ap.add_argument("--csv", default=str(DEFAULT_CSV))
    ap.add_argument("--ferguson", help="owner's Ferguson price sheet (.xlsx); caps prices that would be above Ferguson")
    ap.add_argument("--under", type=float, default=0.10, help="how far below Ferguson to price a capped SKU")
    ap.add_argument("--min-margin", type=float, default=0.20, help="lowest gross margin a Ferguson cap may go to")
    ap.add_argument("--backend", default=os.environ.get("MEDUSA_BACKEND_URL", "https://backend-production-04a8.up.railway.app"))
    ap.add_argument("--env", help="backend .env with MEDUSA_ADMIN_EMAIL / MEDUSA_ADMIN_PASSWORD")
    a = ap.parse_args()
    admin = Admin(a.backend, a.env)

    if a.revert:
        backup = json.load(open(a.revert))
        if backup["backend"] != admin.base:
            sys.exit(f"backup is for {backup['backend']}, not {admin.base}")
        by_product = {}
        for v in backup["variants"]:
            by_product.setdefault(v["product_id"], []).append({"id": v["id"], "prices": price_payload(v["prices"])})
        print("restored", send(admin, by_product), "variants from", a.revert)
        return
    if not a.mult:
        sys.exit("--mult is required")

    cost, notes = {}, {}
    for r in csv.DictReader(open(a.csv, encoding="utf-8-sig")):
        sku = r["SKU"].strip()
        cost[sku] = float(r["Correct Price"] or 0)
        if (r.get("Notes") or "OK").strip() != "OK":
            notes[sku] = r["Notes"].strip()

    ferg = ferguson_prices(a.ferguson) if a.ferguson else {}
    capped = above = 0
    products = admin.products()
    changes, backup, skipped = {}, [], {"tube": 0, "not in cost file": [], "zero cost": [], "flagged in cost file": [], "no plain USD price": [], "unchanged": 0}
    up = down = 0
    print(f"{'SKU':22} {'product':40} {'cost':>9} {'old':>10} {'new':>10}")
    for p in products:
        for v in p.get("variants") or []:
            sku = (v.get("sku") or "").strip()
            if not sku:
                continue
            if is_tube(sku):
                skipped["tube"] += 1; continue
            if sku in notes:
                skipped["flagged in cost file"].append(f"{sku}: {notes[sku]}"); continue
            if sku not in cost:
                skipped["not in cost file"].append(sku); continue
            if cost[sku] <= 0:
                skipped["zero cost"].append(sku); continue
            prices = v.get("prices") or []
            usd = plain_usd(prices)
            if not usd:
                skipped["no plain USD price"].append(sku); continue
            new = target_price(cost[sku], a.mult, ferg.get(sku), a.under, a.min_margin)
            if new < round(cost[sku] * a.mult, 2):
                capped += 1; above += new > ferg[sku]
            if abs(float(usd["amount"]) - new) < 0.005:
                skipped["unchanged"] += 1; continue
            up += new > float(usd["amount"]); down += new < float(usd["amount"])
            print(f"{sku:22} {p['title'][:40]:40} {cost[sku]:9.2f} {float(usd['amount']):10.2f} {new:10.2f}")
            changes.setdefault(p["id"], []).append({"id": v["id"], "prices": price_payload(prices, new)})
            backup.append({"product_id": p["id"], "id": v["id"], "sku": sku, "prices": prices})

    n = sum(len(x) for x in changes.values())
    print(f"\n{n} variants to reprice at cost x {a.mult} ({down} down, {up} up); {skipped['unchanged']} already right; {skipped['tube']} tube skipped")
    if ferg:
        print(f"Ferguson cap: {capped} SKUs priced below cost x {a.mult}; {above} of them stay above Ferguson at the {a.min_margin:.0%} margin floor")
    for k in ("flagged in cost file", "zero cost", "not in cost file", "no plain USD price"):
        if skipped[k]:
            print(f"left alone, {k} ({len(skipped[k])}): " + ", ".join(skipped[k][:40]) + (" ..." if len(skipped[k]) > 40 else ""))
    if not a.apply:
        print("dry run only — add --apply to write"); return
    f = f"pricing-backup-{time.strftime('%Y%m%d-%H%M%S')}.json"
    json.dump({"backend": admin.base, "mult": a.mult, "variants": backup}, open(f, "w"), indent=1)
    print(f"backup of old prices: {f} (undo with --revert {f})")
    print("updated", send(admin, changes), "variants")


if __name__ == "__main__":
    main()
